#!/usr/bin/env python3
"""build_docx_submission.py

Converts verified V2.2 Markdown files into publication-formatted Word (.docx) documents
for JGIM / NMCD journal submission:
- Times New Roman 11/12pt typography
- Clean three-line academic tables (Booktabs style)
- Proper subscripts and superscripts (RERI_PR, AP_PR, kg/m², 10⁹/L, etc.)
- Clean Unicode mathematical typography (≥, ≤, ×, ΔAUC, etc.)
- Zero raw LaTeX $ or Markdown ** syntax in Word output
- Verification of non-empty files and compliance
"""

import os
import re
import shutil
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
SUBMISSION_DIR = BASE / 'submission_package'
RESULTS_DIR = BASE / 'results'

# ---------------------------------------------------------------------------
# TYPOGRAPHY & TEXT PROCESSING
# ---------------------------------------------------------------------------

def clean_math_and_markdown(text):
    """Converts LaTeX math, units, and markdown formatting into tagged text."""
    # 1. Display equations: $$\text{hs-CRP}_{\text{bridged}} = \max(...)$$
    def rep_disp(m):
        eq = m.group(1)
        eq = re.sub(r'\\text\{([^}]+)\}', r'\1', eq)
        eq = re.sub(r'\\max', 'max', eq)
        eq = re.sub(r'\\min', 'min', eq)
        eq = re.sub(r'\\times', '×', eq)
        eq = re.sub(r'_\{([^}]+)\}', r'~sub~\1~/sub~', eq)
        eq = re.sub(r'\^\{([^}]+)\}', r'~sup~\1~/sup~', eq)
        return eq
    text = re.sub(r'\$\$(.*?)\$\$', rep_disp, text, flags=re.DOTALL)

    # 2. Inline math: $...$
    def rep_inline(m):
        s = m.group(1)
        s = re.sub(r'\\ge\b|\\geq\b', '≥', s)
        s = re.sub(r'\\le\b|\\leq\b', '≤', s)
        s = re.sub(r'\\times\b', '×', s)
        s = re.sub(r'\\pm\b', '±', s)
        s = re.sub(r'\\Delta\b', 'Δ', s)
        s = re.sub(r'\\approx\b', '≈', s)
        s = re.sub(r'\\sim\b', '~', s)
        s = re.sub(r'\\%|\\%', '%', s)
        s = re.sub(r'\\text\{([^}]+)\}', r'\1', s)
        s = re.sub(r'\\mathrm\{([^}]+)\}', r'\1', s)
        s = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', s)
        s = re.sub(r'_\{([^}]+)\}', r'~sub~\1~/sub~', s)
        s = re.sub(r'\^\{([^}]+)\}', r'~sup~\1~/sup~', s)
        return s
    text = re.sub(r'\$([^$\n]+)\$', rep_inline, text)

    # 3. Clean any residual LaTeX commands or escaped characters
    text = re.sub(r'\\ge\b|\\geq\b', '≥', text)
    text = re.sub(r'\\le\b|\\leq\b', '≤', text)
    text = re.sub(r'\\times\b', '×', text)
    text = re.sub(r'\\Delta\b', 'Δ', text)
    text = re.sub(r'\\pm\b', '±', text)
    text = re.sub(r'\\%|\\%', '%', text)
    text = re.sub(r'\\([a-zA-Z]+)', r'\1', text)  # remove stray backslashes before words

    # 4. Standardize subscripts for epidemiological interaction terms & odds ratios
    text = re.sub(r'\b(RERI|AP)_(PR|OR)\b', r'\1~sub~\2~/sub~', text)
    text = re.sub(r'\bOR_(11|10|01)\b', r'OR~sub~\1~/sub~', text)
    text = text.replace('OR₁₁', 'OR~sub~11~/sub~')
    text = text.replace('OR₁₀', 'OR~sub~10~/sub~')
    text = text.replace('OR₀₁', 'OR~sub~01~/sub~')
    text = text.replace('hs-CRP_bridged', 'hs-CRP~sub~bridged~/sub~')

    # 5. Standardize superscripts for medical units
    text = text.replace('kg/m2', 'kg/m~sup~2~/sup~')
    text = text.replace('kg/m²', 'kg/m~sup~2~/sup~')
    text = text.replace('10^9/L', '10~sup~9~/sup~/L')
    text = text.replace('10⁹/L', '10~sup~9~/sup~/L')
    text = text.replace('mL/min/1.73m2', 'mL/min/1.73m~sup~2~/sup~')
    text = text.replace('mL/min/1.73m²', 'mL/min/1.73m~sup~2~/sup~')
    text = text.replace('m²', 'm~sup~2~/sup~')

    # 6. Markdown links: [text](url) -> text
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)

    # 7. Markdown bold and italics
    text = re.sub(r'\*\*\*([^*]+)\*\*\*', r'~b~~i~\1~/i~~/b~', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'~b~\1~/b~', text)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'~i~\1~/i~', text)
    text = re.sub(r'`([^`]+)`', r'~code~\1~/code~', text)

    # 8. Ensure no lingering raw $ remain
    text = text.replace('$', '')

    return text


def add_formatted_runs(paragraph, text, base_font_name='Times New Roman', base_font_size=12, is_bold=False, is_italic=False, font_color=None):
    """Parses tagged text and appends styled runs to a python-docx paragraph."""
    tagged = clean_math_and_markdown(text)
    pattern = re.compile(r'(~b~|~/b~|~i~|~/i~|~sub~|~/sub~|~sup~|~/sup~|~code~|~/code~)')
    tokens = pattern.split(tagged)

    cur_bold = is_bold
    cur_italic = is_italic
    cur_sub = False
    cur_sup = False

    for t in tokens:
        if not t:
            continue
        if t == '~b~':
            cur_bold = True
        elif t == '~/b~':
            cur_bold = is_bold
        elif t == '~i~':
            cur_italic = True
        elif t == '~/i~':
            cur_italic = is_italic
        elif t == '~sub~':
            cur_sub = True
        elif t == '~/sub~':
            cur_sub = False
        elif t == '~sup~':
            cur_sup = True
        elif t == '~/sup~':
            cur_sup = False
        elif t == '~code~' or t == '~/code~':
            pass
        else:
            run = paragraph.add_run(t)
            run.font.name = base_font_name
            run.font.size = Pt(base_font_size)
            run.bold = cur_bold
            run.italic = cur_italic
            run.font.subscript = cur_sub
            run.font.superscript = cur_sup
            if font_color:
                run.font.color.rgb = font_color

# ---------------------------------------------------------------------------
# TABLE FORMATTING & STYLING (BOOKTABS THREE-LINE STYLE)
# ---------------------------------------------------------------------------

def apply_booktabs_borders(table):
    """Applies clean academic three-line borders (top 1.5pt, header bottom 1.0pt, bottom 1.5pt)."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')

    # Top border (1.5pt = sz 12)
    top_b = OxmlElement('w:top')
    top_b.set(qn('w:val'), 'single')
    top_b.set(qn('w:sz'), '12')
    top_b.set(qn('w:space'), '0')
    top_b.set(qn('w:color'), '000000')
    tblBorders.append(top_b)

    # Bottom border (1.5pt = sz 12)
    bot_b = OxmlElement('w:bottom')
    bot_b.set(qn('w:val'), 'single')
    bot_b.set(qn('w:sz'), '12')
    bot_b.set(qn('w:space'), '0')
    bot_b.set(qn('w:color'), '000000')
    tblBorders.append(bot_b)

    # No left, right, insideV, insideH borders
    for edge in ['left', 'right', 'insideV', 'insideH']:
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'none')
        tblBorders.append(b)

    tblPr.append(tblBorders)

    # Header bottom border on each cell of row 0
    if len(table.rows) > 0:
        header_tr = table.rows[0]._tr
        trPr = header_tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:tblHeader'))

        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            b = OxmlElement('w:bottom')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '8')  # 1.0pt
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), '000000')
            tcBorders.append(b)
            for edge in ['top', 'left', 'right']:
                nb = OxmlElement(f'w:{edge}')
                nb.set(qn('w:val'), 'none')
                tcBorders.append(nb)
            tcPr.append(tcBorders)

    # Row properties: cantSplit and cell padding
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = OxmlElement('w:tcMar')
            for m, val in [('top', '70'), ('bottom', '70'), ('left', '110'), ('right', '110')]:
                node = OxmlElement(f'w:{m}')
                node.set(qn('w:w'), val)
                node.set(qn('w:type'), 'dxa')
                tcMar.append(node)
            tcPr.append(tcMar)


def add_markdown_table(doc, table_text, col_widths=None, font_size=9.0):
    """Parses a markdown table string and inserts a styled Word table."""
    lines = [l.strip() for l in table_text.strip().splitlines() if l.strip()]
    if len(lines) < 2:
        return

    # Header
    raw_headers = [c.strip() for c in lines[0].strip('|').split('|')]
    start_row = 1
    if len(lines) > 1 and re.match(r'^\|?(\s*:?-+:?\s*\|?)+$', lines[1]):
        start_row = 2

    raw_rows = []
    for line in lines[start_row:]:
        if not line.startswith('|'):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        if len(cells) < len(raw_headers):
            cells += [''] * (len(raw_headers) - len(cells))
        elif len(cells) > len(raw_headers):
            cells = cells[:len(raw_headers)]
        raw_rows.append(cells)

    num_cols = len(raw_headers)
    num_rows = len(raw_rows) + 1

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header row
    for col_idx, h_text in enumerate(raw_headers):
        cell = table.rows[0].cells[col_idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        add_formatted_runs(p, h_text, base_font_size=font_size, is_bold=True)

    # Data rows
    for row_idx, r_cells in enumerate(raw_rows):
        table_row = table.rows[row_idx + 1]
        for col_idx, c_text in enumerate(r_cells):
            cell = table_row.cells[col_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            add_formatted_runs(p, c_text, base_font_size=font_size)

    # Apply widths if specified
    if col_widths and len(col_widths) == num_cols:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    apply_booktabs_borders(table)

    # Add small spacing after table
    after_p = doc.add_paragraph()
    after_p.paragraph_format.space_before = Pt(0)
    after_p.paragraph_format.space_after = Pt(6)

# ---------------------------------------------------------------------------
# DOCUMENT PARSER ENGINE
# ---------------------------------------------------------------------------

def parse_markdown_content(doc, md_text, default_font_size=12, table_widths_map=None):
    """Parses complete markdown text into headings, paragraphs, lists, and tables."""
    lines = md_text.splitlines()
    idx = 0
    n = len(lines)

    while idx < n:
        line = lines[idx]
        stripped = line.strip()

        # Blank line
        if not stripped:
            idx += 1
            continue

        # Horizontal rule
        if stripped in ['---', '***', '___']:
            idx += 1
            continue

        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, stripped[2:], base_font_size=16, is_bold=True)
            idx += 1
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, stripped[3:], base_font_size=13, is_bold=True)
            idx += 1
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, stripped[4:], base_font_size=12, is_bold=True)
            idx += 1
            continue

        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, stripped[5:], base_font_size=11, is_bold=True, is_italic=True)
            idx += 1
            continue

        # Table block
        if stripped.startswith('|'):
            table_lines = []
            while idx < n and lines[idx].strip().startswith('|'):
                table_lines.append(lines[idx].strip())
                idx += 1
            table_str = '\n'.join(table_lines)
            widths = None
            if table_widths_map:
                # heuristic match based on header
                for k, w in table_widths_map.items():
                    if k.lower() in table_lines[0].lower():
                        widths = w
                        break
            add_markdown_table(doc, table_str, col_widths=widths, font_size=9.0)
            continue

        # Bullet list
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            # add bullet character run
            b_run = p.add_run('• ')
            b_run.font.name = 'Times New Roman'
            b_run.font.size = Pt(default_font_size)
            add_formatted_runs(p, stripped[2:], base_font_size=default_font_size)
            idx += 1
            continue

        # Numbered list: 1. Item
        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m_num:
            num_str, rest = m_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            n_run = p.add_run(f'{num_str}. ')
            n_run.font.name = 'Times New Roman'
            n_run.font.size = Pt(default_font_size)
            add_formatted_runs(p, rest, base_font_size=default_font_size)
            idx += 1
            continue

        # Display math line
        if stripped.startswith('$$') and stripped.endswith('$$'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_runs(p, stripped, base_font_size=default_font_size, is_italic=True)
            idx += 1
            continue

        # Normal body paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, stripped, base_font_size=default_font_size)
        idx += 1

# ---------------------------------------------------------------------------
# DOCUMENT BUILDERS
# ---------------------------------------------------------------------------

def init_docx(landscape=False):
    """Initializes a clean Document with Times New Roman and standard margins."""
    doc = docx.Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Inches(11.0)
        sec.page_height = Inches(8.5)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)

    # Base style
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def add_landscape_section(doc):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11.0)
    sec.page_height = Inches(8.5)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    return sec


def add_portrait_section(doc):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    return sec

# ---------------------------------------------------------------------------
# BUILD TARGETS
# ---------------------------------------------------------------------------

def build_01_single_file_manuscript():
    """1. 01_Manuscript_Complete_Single_File_JGIM.docx:
    Combines Main Text + Tables 1-3 + Figures 1-2 with page breaks."""
    print('Building 01_Manuscript_Complete_Single_File_JGIM.docx...')
    doc = init_docx(landscape=False)

    # 1. Main text
    ms_text = (BASE / 'MANUSCRIPT_FULL_JGIM_NMCD.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, ms_text, default_font_size=12)

    # 2. Tables Section (Landscape for clear presentation of 6-column Table 1)
    add_landscape_section(doc)

    t1_text = (BASE / 'Table_1_Baseline.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, t1_text, default_font_size=11, table_widths_map={'characteristic': [2.5, 1.4, 1.4, 1.4, 1.4, 1.4]})

    doc.add_page_break()
    t2_text = (BASE / 'Table_2_Multivariable_Joint_Models.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, t2_text, default_font_size=11, table_widths_map={'parameter': [2.9, 2.2, 2.2, 2.2]})

    doc.add_page_break()
    t3_text = (BASE / 'Table_3_VCTE_vs_FIB4_Diagnostic_Comparison.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, t3_text, default_font_size=11, table_widths_map={
        'diagnostic metric': [2.3, 1.8, 1.8, 1.8, 1.8],
        'model specification': [3.5, 2.0, 2.0, 2.0]
    })

    # 3. Figures Section (Portrait)
    add_portrait_section(doc)

    # Heading
    p_fig_head = doc.add_paragraph()
    p_fig_head.paragraph_format.space_before = Pt(12)
    p_fig_head.paragraph_format.space_after = Pt(12)
    add_formatted_runs(p_fig_head, 'FIGURES AND ILLUSTRATIONS', base_font_size=14, is_bold=True)

    # Figure 1
    fig1_path = RESULTS_DIR / 'figure_1_vcte_crp_forest.png'
    if fig1_path.exists():
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.space_before = Pt(6)
        p_img1.paragraph_format.space_after = Pt(6)
        doc.add_picture(str(fig1_path), width=Inches(6.5))

        p_leg1 = doc.add_paragraph()
        p_leg1.paragraph_format.space_before = Pt(6)
        p_leg1.paragraph_format.space_after = Pt(12)
        p_leg1.paragraph_format.line_spacing = 1.15
        leg1_text = (
            "**Figure 1. Joint associations of liver stiffness and inflammation with advanced CKM.** "
            "Panel A shows Model 2 adjusted odds ratios for the four VCTE/hs-CRP phenotypes relative to neither elevation. "
            "Thresholds are VCTE ≥8.0 kPa and hs-CRP ≥2.0 mg/L. Panel B shows dual-high versus reference associations across "
            "Models 1–3, secondary outcomes, and subgroups; secondary outcomes and subgroups use Model 2. Points are odds ratios; "
            "horizontal bars are 95% confidence intervals from full-universe survey Taylor linearization (df = 40). "
            "Axes are logarithmic. Model 2 adjusts for demographics, survey cycle, socioeconomic characteristics, and adiposity. "
            "RERI and AP describe odds-scale interaction; the standardized prevalence contrast is in percentage points. "
            "None is a causal attributable fraction. CKM, cardiovascular-kidney-metabolic syndrome; VCTE, vibration-controlled "
            "transient elastography; hs-CRP, high-sensitivity C-reactive protein; CVD, cardiovascular disease; CKD, chronic kidney "
            "disease; KDIGO, Kidney Disease: Improving Global Outcomes; OR, odds ratio; CI, confidence interval; RERI, relative "
            "excess risk due to interaction; AP, attributable proportion; df, degrees of freedom."
        )
        add_formatted_runs(p_leg1, leg1_text, base_font_size=10.5)

    doc.add_page_break()

    # Figure 2
    fig2_path = RESULTS_DIR / 'figure_2_vcte_vs_fib4_diagnostic.png'
    if fig2_path.exists():
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(6)
        p_img2.paragraph_format.space_after = Pt(6)
        doc.add_picture(str(fig2_path), width=Inches(6.5))

        p_leg2 = doc.add_paragraph()
        p_leg2.paragraph_format.space_before = Pt(6)
        p_leg2.paragraph_format.space_after = Pt(12)
        p_leg2.paragraph_format.line_spacing = 1.15
        leg2_text = (
            "**Figure 2. FIB-4 performance relative to VCTE and apparent CKM discrimination.** "
            "Panel A reports survey-weighted false-positive rate (FPR), specificity, and positive predictive value (PPV) "
            "for FIB-4 ≥1.30 in younger (30–64), older (65–79), and all adults (30–79), and FIB-4 ≥2.00 in older adults. "
            "The reference is VCTE ≥8.0 kPa, not biopsy-confirmed fibrosis. The callout reports the sensitivity trade-off "
            "with age calibration. Panel B plots subsampled empirical survey-weighted receiver operating characteristic (ROC) "
            "vertices for three Model 2 biomarker extensions; tied predictions enter together. AUCs use the full predictions, "
            "not the plotted subsample. The dotted diagonal denotes chance. AUC confidence intervals in Table 3 use a stratified "
            "delete-1-primary-sampling-unit jackknife (df = 40), conditional on fixed full-sample predictions; performance is "
            "in-sample, without external validation. FIB-4, Fibrosis-4 index; VCTE, vibration-controlled transient elastography; "
            "hs-CRP, high-sensitivity C-reactive protein; CKM, cardiovascular-kidney-metabolic syndrome; AUC, area under the "
            "ROC curve; df, degrees of freedom."
        )
        add_formatted_runs(p_leg2, leg2_text, base_font_size=10.5)

    out_file = SUBMISSION_DIR / '01_Manuscript_Complete_Single_File_JGIM.docx'
    doc.save(str(out_file))
    print(f'  Saved: {out_file.name} ({out_file.stat().st_size:,} bytes)')


def build_02_manuscript_main_text():
    """2. 02_Manuscript_Main_Text_Only.docx: MANUSCRIPT_FULL_JGIM_NMCD.md alone."""
    print('Building 02_Manuscript_Main_Text_Only.docx...')
    doc = init_docx(landscape=False)
    ms_text = (BASE / 'MANUSCRIPT_FULL_JGIM_NMCD.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, ms_text, default_font_size=12)
    out_file = SUBMISSION_DIR / '02_Manuscript_Main_Text_Only.docx'
    doc.save(str(out_file))
    print(f'  Saved: {out_file.name} ({out_file.stat().st_size:,} bytes)')


def build_03_tables_document():
    """3. 03_Tables_1_to_3.docx: Tables 1, 2, and 3 on separate pages."""
    print('Building 03_Tables_1_to_3.docx...')
    doc = init_docx(landscape=True)  # Landscape accommodates 6-column Table 1 beautifully

    t1_text = (BASE / 'Table_1_Baseline.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, t1_text, default_font_size=11, table_widths_map={'characteristic': [2.5, 1.4, 1.4, 1.4, 1.4, 1.4]})

    doc.add_page_break()
    t2_text = (BASE / 'Table_2_Multivariable_Joint_Models.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, t2_text, default_font_size=11, table_widths_map={'parameter': [2.9, 2.2, 2.2, 2.2]})

    doc.add_page_break()
    t3_text = (BASE / 'Table_3_VCTE_vs_FIB4_Diagnostic_Comparison.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, t3_text, default_font_size=11, table_widths_map={
        'diagnostic metric': [2.3, 1.8, 1.8, 1.8, 1.8],
        'model specification': [3.5, 2.0, 2.0, 2.0]
    })

    out_file = SUBMISSION_DIR / '03_Tables_1_to_3.docx'
    doc.save(str(out_file))
    print(f'  Saved: {out_file.name} ({out_file.stat().st_size:,} bytes)')


def build_04_strobe_waterfall():
    """4. 04_Supplementary_STROBE_Waterfall.docx: STROBE_Waterfall_v2.md."""
    print('Building 04_Supplementary_STROBE_Waterfall.docx...')
    doc = init_docx(landscape=True)  # Landscape provides ample room for 5-column waterfall steps
    sw_text = (BASE / 'STROBE_Waterfall_v2.md').read_text(encoding='utf-8')
    parse_markdown_content(doc, sw_text, default_font_size=11, table_widths_map={
        'sequential step': [1.8, 3.5, 1.4, 1.4, 1.4]
    })
    out_file = SUBMISSION_DIR / '04_Supplementary_STROBE_Waterfall.docx'
    doc.save(str(out_file))
    print(f'  Saved: {out_file.name} ({out_file.stat().st_size:,} bytes)')


def build_05_cover_letter():
    """5. 05_Cover_Letter_JGIM.docx: Section 1 of COVER_LETTER_AND_ZERO_FEE_SUBMISSION_GUIDE.md."""
    print('Building 05_Cover_Letter_JGIM.docx...')
    doc = init_docx(landscape=False)
    cl_full = (BASE / 'COVER_LETTER_AND_ZERO_FEE_SUBMISSION_GUIDE.md').read_text(encoding='utf-8')
    sec1 = cl_full.split('## SECTION 1:')[1].split('## SECTION 2:')[0].strip()

    # Clean header prefix if present
    sec1_clean = re.sub(r'^OFFICIAL COVER LETTER FOR EDITORIAL SUBMISSION\s*', '', sec1)
    parse_markdown_content(doc, sec1_clean, default_font_size=11.5)

    out_file = SUBMISSION_DIR / '05_Cover_Letter_JGIM.docx'
    doc.save(str(out_file))
    print(f'  Saved: {out_file.name} ({out_file.stat().st_size:,} bytes)')


def build_06_icmje_disclosure_forms():
    """6. Generates 06_ICMJE_Disclosure_Forms_All_Authors.docx and 06_ICMJE_Disclosure_Form_Gambhirrao.docx
    covering all 3 authors (Dhruv Gambhirrao, Devpal Ghuman, Tarun Achar) with 13-item 'None' tables
    and formal electronic signature blocks dated October 16, 2026."""
    print('Building 06_ICMJE_Disclosure_Forms_All_Authors.docx...')
    doc = init_docx(landscape=False)

    authors = [
        {
            "name": "Dhruv Gambhirrao",
            "affiliation": "Department of Biological Sciences, University of North Texas, Denton, TX, USA",
            "email": "dhruvgambhirrao@my.unt.edu",
            "role": "First Author, Corresponding Author, Principal Investigator",
            "date": "October 16, 2026"
        },
        {
            "name": "Devpal Ghuman",
            "affiliation": "Department of Biological Sciences, University of North Texas, Denton, TX, USA",
            "email": "dsg3583@gmail.com",
            "role": "Co-Author, Epidemiological Investigator",
            "date": "October 16, 2026"
        },
        {
            "name": "Tarun Achar",
            "affiliation": "Department of Biological Sciences, University of North Texas, Denton, TX, USA",
            "email": "tarunah.achar@gmail.com",
            "role": "Co-Author, Clinical Research Investigator",
            "date": "October 16, 2026"
        }
    ]

    domains = [
        ("1. Grants or contracts from any entity", "None", "No external research grant or contract support received."),
        ("2. Royalties or licenses", "None", "No royalties, licensing fees, or intellectual property income."),
        ("3. Consulting fees", "None", "No consulting fees, advisory board payments, or retainer income."),
        ("4. Payment or honoraria for lectures, presentations, speakers bureaus, manuscript writing, or educational events", "None", "No speaking honoraria, presentation fees, or educational stipends."),
        ("5. Payment for expert testimony", "None", "No expert witness or medico-legal testimony fees."),
        ("6. Support for attending meetings and/or travel", "None", "No travel grants, hotel accommodations, or conference registration support."),
        ("7. Patents planned, issued, or pending", "None", "No patents planned, pending, or licensed related to the work."),
        ("8. Participation on a Data Safety Monitoring Board (DSMB) or Advisory Board", "None", "No participation or compensation from any clinical trial DSMB or advisory board."),
        ("9. Leadership or fiduciary role in other board, society, committee, or advocacy group, paid or unpaid", "None", "No officer, board member, committee chair, or fiduciary governance positions."),
        ("10. Stock or stock options", "None", "No equity ownership, stock holdings, or options in pharmaceutical, diagnostic, or medical device corporations."),
        ("11. Receipt of equipment, materials, drugs, medical writing, gifts, or other services", "None", "No hardware, software, reagent donations, medical writing assistance, or clinical gifts."),
        ("12. Other financial or non-financial interests", "None", "No competing non-financial relationships, professional rivalries, or personal conflicts."),
        ("13. All support for the present manuscript (funding, materials, APC subsidies, or sponsored data analysis)", "None", "The study was conducted independently without external funding; all NHANES datasets are public-domain files.")
    ]

    for auth_idx, auth in enumerate(authors):
        # Header
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(2)
        r_title = p_title.add_run("ICMJE DISCLOSURE FORM")
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(15)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(12)
        r_sub = p_sub.add_run("International Committee of Medical Journal Editors (ICMJE)\nDisclosure of Potential Conflicts of Interest")
        r_sub.font.name = "Times New Roman"
        r_sub.font.size = Pt(10.5)
        r_sub.italic = True

        # Section 1
        p_s1 = doc.add_paragraph()
        p_s1.paragraph_format.space_before = Pt(4)
        p_s1.paragraph_format.space_after = Pt(4)
        r_s1 = p_s1.add_run("SECTION 1: IDENTIFYING INFORMATION")
        r_s1.bold = True
        r_s1.font.name = "Times New Roman"
        r_s1.font.size = Pt(11.5)

        meta_items = [
            ("Author Full Name:", auth["name"]),
            ("Institutional Affiliation:", auth["affiliation"]),
            ("Author Email Address:", auth["email"]),
            ("Author Role in Study:", auth["role"]),
            ("Date of Completion:", auth["date"]),
            ("Manuscript Title:", (
                "Direct Physical Vibration-Controlled Transient Elastography vs. Formula-Based FIB-4 and "
                "High-Sensitivity C-Reactive Protein in Advanced Cardiovascular-Kidney-Metabolic Syndrome: "
                "A Nationally Representative Complex Survey Cohort Study (NHANES 2017–2023)"
            )),
            ("Target Journals:", "Journal of General Internal Medicine (JGIM) / Nutrition, Metabolism and Cardiovascular Diseases (NMCD)")
        ]

        for label, val in meta_items:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.10
            r_lbl = p.add_run(f"{label} ")
            r_lbl.bold = True
            r_lbl.font.name = "Times New Roman"
            r_lbl.font.size = Pt(10)
            r_val = p.add_run(val)
            r_val.font.name = "Times New Roman"
            r_val.font.size = Pt(10)

        # Section 2
        p_s2 = doc.add_paragraph()
        p_s2.paragraph_format.space_before = Pt(8)
        p_s2.paragraph_format.space_after = Pt(4)
        r_s2 = p_s2.add_run("SECTION 2: DISCLOSURE OF INTERESTS (13 STANDARD ICMJE DOMAINS)")
        r_s2.bold = True
        r_s2.font.name = "Times New Roman"
        r_s2.font.size = Pt(11.5)

        # Table
        table = doc.add_table(rows=len(domains) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ["ICMJE Disclosure Domain", "Status", "Details / Explanation"]
        col_widths = [Inches(2.8), Inches(0.85), Inches(2.85)]

        for c_idx, h_text in enumerate(headers):
            cell = table.rows[0].cells[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(h_text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)

        for r_idx, (dom_name, status, details) in enumerate(domains):
            row = table.rows[r_idx + 1]
            p0 = row.cells[0].paragraphs[0]
            p0.paragraph_format.space_before = Pt(1.5)
            p0.paragraph_format.space_after = Pt(1.5)
            r0 = p0.add_run(dom_name)
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(9)

            p1 = row.cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_before = Pt(1.5)
            p1.paragraph_format.space_after = Pt(1.5)
            r1 = p1.add_run(status)
            r1.bold = True
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9)

            p2 = row.cells[2].paragraphs[0]
            p2.paragraph_format.space_before = Pt(1.5)
            p2.paragraph_format.space_after = Pt(1.5)
            r2 = p2.add_run(details)
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(9)

        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = w

        apply_booktabs_borders(table)

        # Section 3
        p_s3 = doc.add_paragraph()
        p_s3.paragraph_format.space_before = Pt(8)
        p_s3.paragraph_format.space_after = Pt(3)
        r_s3 = p_s3.add_run("SECTION 3: FORMAL ATTESTATION AND CERTIFICATION")
        r_s3.bold = True
        r_s3.font.name = "Times New Roman"
        r_s3.font.size = Pt(11.5)

        p_cert = doc.add_paragraph()
        p_cert.paragraph_format.space_before = Pt(0)
        p_cert.paragraph_format.space_after = Pt(6)
        p_cert.paragraph_format.line_spacing = 1.10
        r_cert = p_cert.add_run(
            "I certify that I have read the ICMJE Form for Disclosure of Potential Conflicts of Interest, "
            "that the information provided herein is accurate, truthful, and complete to the best of my knowledge, "
            "and that I have disclosed all relevant relationships within the past 36 months as well as all specific "
            "support for the submitted manuscript."
        )
        r_cert.font.name = "Times New Roman"
        r_cert.font.size = Pt(9.5)

        # Signature Table
        sig_table = doc.add_table(rows=3, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = False
        sig_widths = [Inches(3.25), Inches(3.25)]

        sig_rows_content = [
            [("Electronic Signature: ", True), (auth["name"], False)],
            [("Author Full Name: ", True), (auth["name"], False)],
            [("Date of Signature: ", True), (auth["date"], False)]
        ]

        for r_i, content in enumerate(sig_rows_content):
            row = sig_table.rows[r_i]
            p0 = row.cells[0].paragraphs[0]
            p0.paragraph_format.space_before = Pt(2)
            p0.paragraph_format.space_after = Pt(2)
            r0 = p0.add_run(content[0][0])
            r0.bold = True
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(9.5)

            p1 = row.cells[1].paragraphs[0]
            p1.paragraph_format.space_before = Pt(2)
            p1.paragraph_format.space_after = Pt(2)
            r1 = p1.add_run(content[1][0])
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9.5)

        for row in sig_table.rows:
            for c_idx, w in enumerate(sig_widths):
                row.cells[c_idx].width = w

        apply_booktabs_borders(sig_table)

        if auth_idx < len(authors) - 1:
            doc.add_page_break()

    out_file_all = SUBMISSION_DIR / '06_ICMJE_Disclosure_Forms_All_Authors.docx'
    doc.save(str(out_file_all))
    print(f'  Saved: {out_file_all.name} ({out_file_all.stat().st_size:,} bytes)')

    out_file_g = SUBMISSION_DIR / '06_ICMJE_Disclosure_Form_Gambhirrao.docx'
    doc.save(str(out_file_g))
def copy_figure_assets():
    """Copies high-res TIFF and PNG figures into submission_package/."""
    print('Copying high-resolution figures to submission_package/...')
    fig_names = [
        'figure_1_vcte_crp_forest.tiff',
        'figure_2_vcte_vs_fib4_diagnostic.tiff',
        'figure_1_vcte_crp_forest.png',
        'figure_2_vcte_vs_fib4_diagnostic.png'
    ]
    for fn in fig_names:
        src = RESULTS_DIR / fn
        dst = SUBMISSION_DIR / fn
        if src.exists():
            shutil.copy2(src, dst)
            print(f'  Copied: {fn} ({dst.stat().st_size:,} bytes)')
        else:
            print(f'  [WARNING] Missing figure source: {src}')


# ---------------------------------------------------------------------------
# VERIFICATION ENGINE
# ---------------------------------------------------------------------------

def verify_submission_package():
    """Scans all generated .docx files to assert non-empty, 0 raw $, and 0 Markdown **."""
    print('\n=== Verifying Submission Package Documents ===')
    docx_files = sorted(SUBMISSION_DIR.glob('*.docx'))
    assert len(docx_files) >= 6, f"Expected at least 6 docx files, found {len(docx_files)}"

    for df in docx_files:
        assert df.stat().st_size > 0, f"File {df.name} is empty!"
        doc = docx.Document(str(df))
        
        # Verify text in paragraphs
        para_texts = [p.text for p in doc.paragraphs]
        table_texts = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        all_text = ' '.join(para_texts + table_texts)

        # Check for raw LaTeX $
        dollar_count = all_text.count('$')
        assert dollar_count == 0, f"Raw $ detected in {df.name} (count={dollar_count})"

        # Check for raw Markdown **
        bold_stars = all_text.count('**')
        assert bold_stars == 0, f"Raw ** detected in {df.name} (count={bold_stars})"

        # Check for raw ### headings
        hash_count = len(re.findall(r'#+\s', all_text))
        assert hash_count == 0, f"Raw Markdown # heading detected in {df.name} (count={hash_count})"

        print(f'  [PASS] {df.name}: {df.stat().st_size:,} bytes | paragraphs: {len(doc.paragraphs)} | tables: {len(doc.tables)} | zero raw $ and zero **')

    # Verify figures
    for fn in ['figure_1_vcte_crp_forest.tiff', 'figure_2_vcte_vs_fib4_diagnostic.tiff',
               'figure_1_vcte_crp_forest.png', 'figure_2_vcte_vs_fib4_diagnostic.png']:
        fp = SUBMISSION_DIR / fn
        assert fp.exists() and fp.stat().st_size > 0, f"Figure {fn} missing or empty"
        print(f'  [PASS] Figure verified: {fn} ({fp.stat().st_size:,} bytes)')

    # Mirror to root directory v2_elastography_manuscript/submission_package if separate
    root_pkg = ROOT / 'v2_elastography_manuscript' / 'submission_package'
    if root_pkg.resolve() != SUBMISSION_DIR.resolve():
        root_pkg.mkdir(parents=True, exist_ok=True)
        for item in SUBMISSION_DIR.iterdir():
            shutil.copy2(item, root_pkg / item.name)
        print(f'  [PASS] Mirrored submission package to {root_pkg}')

    print('\n=== ALL SUBMISSION PACKAGE ASSETS VERIFIED WITH ZERO ERRORS ===')


def main():
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    build_01_single_file_manuscript()
    build_02_manuscript_main_text()
    build_03_tables_document()
    build_04_strobe_waterfall()
    build_05_cover_letter()
    build_06_icmje_disclosure_forms()
    copy_figure_assets()
    verify_submission_package()


if __name__ == '__main__':
    main()
